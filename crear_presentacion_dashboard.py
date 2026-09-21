from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUTPUT = "presentacion_dashboard_congresos.pptx"
WIDE = (Inches(13.333), Inches(7.5))

NAVY = RGBColor(18, 18, 18)
INK = RGBColor(0, 0, 0)
SLATE = RGBColor(92, 92, 92)
MUTED = RGBColor(166, 166, 166)
WHITE = RGBColor(255, 255, 255)
CYAN = RGBColor(242, 126, 32)
BLUE = RGBColor(213, 91, 18)
AMBER = RGBColor(255, 157, 61)
CORAL = RGBColor(180, 65, 12)
PALE = RGBColor(242, 242, 242)
PANEL = RGBColor(38, 38, 38)


def set_bg(slide, color=NAVY):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def text(slide, value, x, y, w, h, size=18, color=WHITE, bold=False,
         font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Pt(0)
    tf.margin_right = Pt(0)
    tf.margin_top = Pt(0)
    tf.margin_bottom = Pt(0)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill, radius=False, line=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def circle(slide, x, y, d, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def line(slide, x1, y1, x2, y2, color=SLATE, width=1.5):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def footer(slide, number):
    text(slide, "DASHBOARD DE CONGRESOS  /  SISTEMA DE INFORMACIÓN ACADÉMICA", 0.65, 7.1, 9.5, 0.18, 7, MUTED, True)
    text(slide, f"{number:02d}", 12.15, 7.05, 0.5, 0.22, 8, CYAN, True, align=PP_ALIGN.RIGHT)


def title(slide, kicker, heading, subtitle=None, number=1):
    text(slide, kicker.upper(), 0.7, 0.55, 6.5, 0.25, 9, CYAN, True)
    text(slide, heading, 0.7, 0.92, 11.6, 0.62, 27, WHITE, True)
    if subtitle:
        text(slide, subtitle, 0.72, 1.68, 10.8, 0.42, 12, MUTED)
    footer(slide, number)


def bullet_list(slide, items, x, y, w, gap=0.53, size=15, color=WHITE, marker=CYAN):
    for index, item in enumerate(items):
        yy = y + index * gap
        circle(slide, x, yy + 0.08, 0.1, marker)
        text(slide, item, x + 0.25, yy, w - 0.25, 0.35, size, color)


def mini_card(slide, x, y, w, h, label, value, accent=CYAN):
    rect(slide, x, y, w, h, PANEL, True, PANEL)
    rect(slide, x, y, 0.08, h, accent)
    text(slide, label.upper(), x + 0.25, y + 0.2, w - 0.4, 0.2, 8, MUTED, True)
    text(slide, value, x + 0.25, y + 0.52, w - 0.4, 0.42, 20, WHITE, True)


def add_slide(prs, bg=NAVY):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg)
    return slide


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = WIDE

    # 01 - Cover
    slide = add_slide(prs)
    rect(slide, 8.65, 0, 4.68, 7.5, RGBColor(17, 47, 68))
    circle(slide, 9.55, 1.0, 2.9, RGBColor(25, 83, 94))
    circle(slide, 10.15, 1.6, 1.7, CYAN)
    circle(slide, 10.67, 2.12, 0.66, NAVY)
    line(slide, 9.5, 4.7, 12.35, 1.85, CYAN, 2)
    line(slide, 9.75, 5.35, 12.4, 4.35, AMBER, 2)
    text(slide, "SISTEMA DE INFORMACIÓN ACADÉMICA", 0.78, 0.82, 6.8, 0.28, 10, CYAN, True)
    text(slide, "Dashboard\nde Congresos", 0.75, 1.45, 7.5, 1.48, 38, WHITE, True)
    text(slide, "Cómo funciona, qué herramientas lo construyen y hacia dónde puede crecer.", 0.8, 3.35, 6.6, 0.72, 17, PALE)
    rect(slide, 0.8, 5.45, 2.25, 0.06, CYAN)
    text(slide, "Presentación del sistema", 0.8, 5.7, 4.5, 0.3, 12, MUTED, True)
    text(slide, "2026", 0.8, 6.22, 1.2, 0.3, 11, WHITE, True)
    text(slide, "01", 12.15, 7.05, 0.5, 0.22, 8, CYAN, True, align=PP_ALIGN.RIGHT)

    # 02 - Purpose
    slide = add_slide(prs)
    title(slide, "01  /  Contexto", "Una sola vista para decidir mejor", "El sistema convierte información académica dispersa en una herramienta de consulta rápida.", 2)
    rect(slide, 0.75, 2.45, 3.7, 3.65, PANEL, True)
    text(slide, "ANTES", 1.05, 2.8, 1.2, 0.25, 9, CORAL, True)
    bullet_list(slide, ["Fuentes separadas", "Búsqueda manual", "Fechas límite difíciles de priorizar", "Poca visibilidad comparativa"], 1.05, 3.3, 3.0, 0.62, 14, PALE, CORAL)
    line(slide, 4.7, 4.25, 5.75, 4.25, CYAN, 2)
    circle(slide, 5.55, 4.08, 0.34, CYAN)
    rect(slide, 6.1, 2.45, 6.45, 3.65, RGBColor(22, 63, 77), True)
    text(slide, "AHORA", 6.45, 2.8, 1.2, 0.25, 9, CYAN, True)
    bullet_list(slide, ["Congresos y revistas centralizados", "Filtros por búsqueda, país y taxonomía", "Estados de deadlines visibles", "Datos listos para actualizar y publicar"], 6.45, 3.3, 5.45, 0.62, 14, WHITE, CYAN)

    # 03 - System overview
    slide = add_slide(prs)
    title(slide, "02  /  El sistema", "Dos módulos, una experiencia", "El usuario parte de un selector central y entra al módulo que necesita.", 3)
    rect(slide, 4.95, 2.2, 3.45, 1.05, RGBColor(22, 63, 77), True)
    text(slide, "SELECTOR DE MÓDULOS", 5.25, 2.47, 2.9, 0.2, 10, CYAN, True, align=PP_ALIGN.CENTER)
    line(slide, 6.65, 3.3, 3.25, 4.25, CYAN, 2)
    line(slide, 6.65, 3.3, 10.05, 4.25, CYAN, 2)
    rect(slide, 1.0, 4.15, 4.5, 1.78, PANEL, True)
    circle(slide, 1.38, 4.58, 0.52, BLUE)
    text(slide, "C", 1.38, 4.62, 0.52, 0.3, 16, NAVY, True, align=PP_ALIGN.CENTER)
    text(slide, "CONGRESOS", 2.1, 4.42, 2.2, 0.26, 15, WHITE, True)
    text(slide, "Oportunidades, fechas límite, países, modalidades y publicación.", 2.1, 4.82, 2.95, 0.48, 11, MUTED)
    rect(slide, 7.82, 4.15, 4.5, 1.78, PANEL, True)
    circle(slide, 8.2, 4.58, 0.52, AMBER)
    text(slide, "R", 8.2, 4.62, 0.52, 0.3, 16, NAVY, True, align=PP_ALIGN.CENTER)
    text(slide, "REVISTAS", 8.92, 4.42, 2.2, 0.26, 15, WHITE, True)
    text(slide, "Consulta de revistas, APC, disciplinas y posibles malas prácticas.", 8.92, 4.82, 2.95, 0.48, 11, MUTED)

    # 04 - Congresses
    slide = add_slide(prs)
    title(slide, "03  /  Módulo Congresos", "Encontrar la oportunidad correcta", "La información se puede explorar por filtros, tarjetas, listado y línea de tiempo.", 4)
    mini_card(slide, 0.78, 2.3, 2.15, 1.12, "Congresos", "1,248", BLUE)
    mini_card(slide, 3.1, 2.3, 2.15, 1.12, "Urgentes", "24", CORAL)
    mini_card(slide, 5.42, 2.3, 2.15, 1.12, "Países", "38", CYAN)
    mini_card(slide, 7.74, 2.3, 2.15, 1.12, "Próximos", "86", AMBER)
    rect(slide, 0.78, 3.8, 4.4, 2.45, PANEL, True)
    text(slide, "FILTROS", 1.08, 4.08, 1.2, 0.2, 9, CYAN, True)
    for i, label in enumerate(["Buscar por nombre o disciplina", "País y modalidad", "Categoría / línea / sublínea", "Indexación y estado"]):
        rect(slide, 1.08, 4.5 + i * 0.36, 3.75, 0.24, RGBColor(39, 55, 76), True)
        text(slide, label, 1.25, 4.54 + i * 0.36, 3.35, 0.16, 9, PALE)
    rect(slide, 5.48, 3.8, 7.05, 2.45, RGBColor(22, 63, 77), True)
    text(slide, "VISTA DE RESULTADOS", 5.8, 4.08, 2.5, 0.2, 9, CYAN, True)
    for i, (name, place, status, color) in enumerate([("World Digital Health", "Madrid  ·  Híbrido", "18 días", AMBER), ("AI & Society Forum", "Toronto  ·  Presencial", "6 días", CORAL), ("Sustainable Cities", "Lisboa  ·  Virtual", "42 días", CYAN)]):
        yy = 4.48 + i * 0.52
        text(slide, name, 5.8, yy, 2.45, 0.2, 11, WHITE, True)
        text(slide, place, 8.3, yy, 2.25, 0.2, 9, MUTED)
        text(slide, status, 11.2, yy, 0.9, 0.2, 10, color, True, align=PP_ALIGN.RIGHT)
        line(slide, 5.8, yy + 0.3, 12.2, yy + 0.3, RGBColor(49, 70, 91), 0.7)

    # 05 - Journals
    slide = add_slide(prs)
    title(slide, "04  /  Módulo Revistas", "Evaluar fuentes con más contexto", "El módulo permite buscar revistas y revisar señales relevantes para publicar o investigar.", 5)
    rect(slide, 0.78, 2.3, 7.2, 3.8, PANEL, True)
    text(slide, "FICHA DE REVISTA", 1.15, 2.65, 2.4, 0.2, 9, AMBER, True)
    text(slide, "Journal of Digital Society", 1.15, 3.05, 5.4, 0.34, 19, WHITE, True)
    text(slide, "ISSN 2049-1234   ·   Editorial académica", 1.15, 3.5, 4.6, 0.2, 10, MUTED)
    rect(slide, 1.15, 4.05, 1.4, 0.44, RGBColor(57, 74, 92), True)
    text(slide, "APC opcional", 1.25, 4.16, 1.2, 0.15, 9, PALE, True, align=PP_ALIGN.CENTER)
    rect(slide, 2.75, 4.05, 1.7, 0.44, RGBColor(24, 79, 76), True)
    text(slide, "Vigente", 2.85, 4.16, 1.5, 0.15, 9, CYAN, True, align=PP_ALIGN.CENTER)
    text(slide, "Clasificación por disciplina y taxonomía institucional", 1.15, 4.85, 5.8, 0.25, 11, PALE)
    rect(slide, 8.35, 2.3, 4.18, 3.8, RGBColor(22, 63, 77), True)
    text(slide, "CONTROLES DE CONSULTA", 8.7, 2.65, 2.8, 0.2, 9, CYAN, True)
    bullet_list(slide, ["Nombre, ISSN o editorial", "Tipo de revista y APC", "Disciplina y categorías", "Revisión de malas prácticas"], 8.7, 3.25, 3.2, 0.58, 13, WHITE, AMBER)

    # 06 - Demo flow
    slide = add_slide(prs)
    title(slide, "05  /  Cómo se usa", "Del objetivo a una decisión en cuatro pasos", "El flujo está pensado para reducir tiempo de búsqueda y hacer visibles las prioridades.", 6)
    steps = [("01", "Elegir módulo", "Congresos o revistas", BLUE), ("02", "Aplicar filtros", "Taxonomía, país, fecha o disciplina", CYAN), ("03", "Comparar", "Tarjetas, lista y estados", AMBER), ("04", "Abrir detalle", "Consultar enlace y tomar acción", CORAL)]
    for i, (num, head, desc, color) in enumerate(steps):
        x = 0.9 + i * 3.05
        circle(slide, x, 2.72, 0.7, color)
        text(slide, num, x, 2.91, 0.7, 0.2, 13, NAVY, True, align=PP_ALIGN.CENTER)
        if i < 3:
            line(slide, x + 0.82, 3.07, x + 2.7, 3.07, RGBColor(55, 84, 103), 1.5)
        text(slide, head, x, 3.75, 2.4, 0.28, 15, WHITE, True)
        text(slide, desc, x, 4.23, 2.35, 0.62, 11, MUTED)
    rect(slide, 0.9, 5.55, 11.45, 0.62, PANEL, True)
    text(slide, "Resultado: una consulta que puede convertirse en seguimiento, postulación o decisión editorial.", 1.2, 5.77, 10.8, 0.2, 12, PALE, True, align=PP_ALIGN.CENTER)

    # 07 - Architecture
    slide = add_slide(prs)
    title(slide, "06  /  Arquitectura", "Una aplicación ligera y fácil de desplegar", "Separar frontend, procesamiento y datos permite actualizar cada parte sin romper la experiencia.", 7)
    rect(slide, 0.85, 2.35, 3.2, 2.65, RGBColor(22, 63, 77), True)
    text(slide, "FRONTEND", 1.18, 2.72, 1.6, 0.2, 10, CYAN, True)
    text(slide, "React + Vite", 1.18, 3.12, 2.3, 0.3, 19, WHITE, True)
    text(slide, "Componentes, filtros, módulos y visualización.", 1.18, 3.72, 2.25, 0.5, 11, PALE)
    rect(slide, 5.05, 2.35, 3.2, 2.65, PANEL, True)
    text(slide, "DATOS", 5.38, 2.72, 1.6, 0.2, 10, AMBER, True)
    text(slide, "JSON procesado", 5.38, 3.12, 2.3, 0.3, 19, WHITE, True)
    text(slide, "Congresos, revistas y taxonomías listas para consumo.", 5.38, 3.72, 2.25, 0.5, 11, PALE)
    rect(slide, 9.25, 2.35, 3.2, 2.65, RGBColor(35, 48, 67), True)
    text(slide, "BACKEND", 9.58, 2.72, 1.6, 0.2, 10, CORAL, True)
    text(slide, "Python", 9.58, 3.12, 2.3, 0.3, 19, WHITE, True)
    text(slide, "Procesamiento de fuentes y automatizaciones.", 9.58, 3.72, 2.25, 0.5, 11, PALE)
    line(slide, 4.12, 3.68, 4.95, 3.68, CYAN, 2)
    line(slide, 8.32, 3.68, 9.15, 3.68, AMBER, 2)
    text(slide, "Render", 5.7, 5.75, 1.8, 0.25, 13, CYAN, True, align=PP_ALIGN.CENTER)
    text(slide, "Despliegue del frontend desde /frontend", 4.35, 6.08, 4.6, 0.22, 10, MUTED, align=PP_ALIGN.CENTER)

    # 08 - Tools
    slide = add_slide(prs)
    title(slide, "07  /  Herramientas", "Tecnologías elegidas para mantener velocidad y claridad", "Cada herramienta cubre una responsabilidad concreta del sistema.", 8)
    tools = [("React", "Interfaz por componentes", BLUE), ("Vite", "Desarrollo y build", CYAN), ("Python", "Procesamiento de datos", AMBER), ("JSON", "Formato de intercambio", CORAL), ("GitHub", "Versionado del proyecto", RGBColor(164, 137, 255)), ("Render", "Despliegue web", RGBColor(122, 205, 226))]
    for i, (name, desc, color) in enumerate(tools):
        col = i % 3
        row = i // 3
        x = 0.85 + col * 4.15
        y = 2.45 + row * 1.55
        rect(slide, x, y, 3.55, 1.12, PANEL, True)
        circle(slide, x + 0.28, y + 0.3, 0.46, color)
        text(slide, name, x + 0.95, y + 0.22, 2.2, 0.25, 15, WHITE, True)
        text(slide, desc, x + 0.95, y + 0.62, 2.25, 0.2, 10, MUTED)

    # 09 - Data/deploy
    slide = add_slide(prs)
    title(slide, "08  /  Operación", "Actualizar datos sin rehacer la aplicación", "El contenido se procesa, se publica como JSON y el frontend lo consume al cargar.", 9)
    stages = [("FUENTE", "Hoja de cálculo\no fuente externa", BLUE), ("PROCESAMIENTO", "Scripts Python\nvalidan y clasifican", AMBER), ("PUBLICACIÓN", "frontend/public\narchivos JSON", CYAN), ("CONSUMO", "React + Render\nexperiencia web", CORAL)]
    for i, (head, body, color) in enumerate(stages):
        x = 0.7 + i * 3.12
        rect(slide, x, 2.65, 2.48, 1.72, PANEL, True)
        circle(slide, x + 0.23, 2.95, 0.42, color)
        text(slide, head, x + 0.82, 2.95, 1.35, 0.2, 9, color, True)
        text(slide, body, x + 0.25, 3.55, 1.95, 0.48, 13, WHITE, True)
        if i < 3:
            line(slide, x + 2.55, 3.5, x + 3.0, 3.5, color, 2)
    rect(slide, 1.15, 5.25, 10.95, 0.72, RGBColor(22, 63, 77), True)
    text(slide, "Configuración de Render: Root Directory = frontend  ·  Build = npm install && npm run build  ·  Publish = dist", 1.45, 5.52, 10.35, 0.18, 10, PALE, True, align=PP_ALIGN.CENTER)

    # 10 - Roadmap
    slide = add_slide(prs)
    title(slide, "09  /  Próximos pasos", "Un sistema preparado para crecer", "Las mejoras futuras priorizan anticipación, colaboración y calidad de datos.", 10)
    roadmap = [("01", "Alertas y recordatorios", "Notificaciones antes de cada deadline.", CORAL), ("02", "Calendario personal", "Exportar eventos a Google Calendar o iCal.", CYAN), ("03", "Analítica avanzada", "Embudo de uso, búsquedas y decisiones.", AMBER), ("04", "Roles y colaboración", "Favoritos, notas y seguimiento por usuario.", BLUE)]
    for i, (num, head, desc, color) in enumerate(roadmap):
        x = 0.85 + (i % 2) * 6.1
        y = 2.35 + (i // 2) * 1.78
        rect(slide, x, y, 5.45, 1.3, PANEL, True)
        text(slide, num, x + 0.3, y + 0.26, 0.48, 0.24, 13, color, True)
        text(slide, head, x + 1.0, y + 0.22, 3.9, 0.26, 15, WHITE, True)
        text(slide, desc, x + 1.0, y + 0.68, 4.0, 0.28, 11, MUTED)
    rect(slide, 0.85, 6.15, 11.45, 0.42, RGBColor(22, 63, 77), True)
    text(slide, "La base actual ya permite incorporar nuevas fuentes y módulos con bajo impacto.", 1.1, 6.28, 10.95, 0.16, 10, PALE, True, align=PP_ALIGN.CENTER)

    prs.save(OUTPUT)
    print(f"Presentación creada: {OUTPUT}")


if __name__ == "__main__":
    build()